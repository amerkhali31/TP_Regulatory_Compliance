import pandas as pd
import os
import constants
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

def mergeBooks(invoice_book_name: str, customer_book_name: str, product_book_name: str,
                invoice_sheet_name: str, customer_sheet_name: str, product_sheet_name: str) -> None:

    '''
    Make One workbook that has all of the sheets from the parent books.
    Gives consistent format to work off of.
    '''
    # Load specific sheets from the source Excel workbooks
    sheet1 = pd.read_excel(invoice_book_name, sheet_name=invoice_sheet_name)
    sheet2 = pd.read_excel(customer_book_name, sheet_name=customer_sheet_name)
    sheet3 = pd.read_excel(product_book_name, sheet_name=product_sheet_name)

    # Create a new Excel writer to save the combined sheets
    with pd.ExcelWriter(f"{constants.DATA_PATH}/aggregated_book.xlsx", engine="openpyxl") as writer:
        # Write each sheet to the new workbook
        sheet1.to_excel(writer, sheet_name=constants.AGGREGATED_BOOK_INVOICE_SHEET_NAME, index=False)
        sheet2.to_excel(writer, sheet_name=constants.AGGREGATED_BOOK_CUSTOMER_SHEET_NAME, index=False)
        sheet3.to_excel(writer, sheet_name=constants.AGGREGATED_BOOK_PRODUCT_SHEET_NAME, index=False)

def clear_unnamed_columns(df: pd.DataFrame):
    df.drop(columns=[col for col in df.columns if 'Unnamed' in col], inplace=True)

def read_sheet(book_name: str, sheet_name: str, cols: list, dtypes: dict = {}) -> pd.DataFrame:
    df = pd.read_excel(book_name, sheet_name=sheet_name, usecols=cols, dtype=dtypes)
    clear_unnamed_columns(df)

    return df

def merge_sheets():
    '''
    Take the sheets in the aggregate book and make the report from it
    '''

    pass


def extract_liability(book_name: str, sheet_name: str, cell: str) -> float:

    # Load the workbook
    wb = load_workbook(book_name, data_only=True)  # data_only=True if you want the *calculated value* of formulas
    ws = wb[sheet_name]  # name of the sheet you're reading from

    return ws[cell].value


def add_filters():
    pass


def autoresize_and_filter(book_name: str, sheet_name: str):
    wb = load_workbook(book_name)
    ws = wb[sheet_name]
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)  # get column letter

        for cell in col:
            if cell.value is not None:
                cell_length = len(str(cell.value))
                if cell_length > max_length:
                    max_length = cell_length

        adjusted_width = max_length + 2  # add some padding
        ws.column_dimensions[col_letter].width = adjusted_width

    # ============================
    # Add filters to all columns
    # ============================
    ws.auto_filter.ref = ws.dimensions  # automatically applies to the used range

    # Save workbook
    wb.save(book_name)

